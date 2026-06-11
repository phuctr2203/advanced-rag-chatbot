#!/usr/bin/env python3
"""Create CSV, JSON, Markdown, and Excel summaries for RAGAS evaluation output."""

from __future__ import annotations

import argparse
import csv
import json
import statistics
import sys
from copy import copy
from collections import defaultdict
from pathlib import Path


SCORE_COLUMNS = [
    "context_precision",
    "context_recall",
    "context_relevance",
    "faithfulness",
    "answer_relevancy",
    "answer_correctness",
]

QUESTION_COLUMNS = [
    "Row ID",
    "Language",
    "Question",
    "Generated Answer",
    "Reference Answer",
    "Reference Context",
    "Expected File",
    "Expected Page",
    "Retrieved Files",
    "Retrieved Pages",
    "Top Retrieved Context",
    "Intent",
    "Detected Language",
    "Retrieval Time Ms",
    "Generation Time Ms",
    "Total Time Ms",
    "Context Precision",
    "Context Recall",
    "Context Relevance",
    "Faithfulness",
    "Answer Relevancy",
    "Answer Correctness",
    "File Hit",
    "Page Hit",
    "Expected Rank",
    "Paraphrase Number",
    "Original Question",
    "Error",
]


def main() -> int:
    args = parse_args()
    input_path = Path(args.input)
    output_dir = Path(args.output_dir) if args.output_dir else input_path.parent
    output_dir.mkdir(parents=True, exist_ok=True)

    records = read_jsonl(input_path)
    rows = [flatten_record(record) for record in records]
    summary = build_summary(rows)

    csv_path = output_dir / "ragas_scores.csv"
    json_path = output_dir / "summary.json"
    markdown_path = output_dir / "summary.md"
    workbook_path = Path(args.workbook) if args.workbook else output_dir / "evaluation_result.xlsx"
    workbook_path.parent.mkdir(parents=True, exist_ok=True)

    write_csv(csv_path, rows)
    json_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    markdown_path.write_text(build_markdown(summary), encoding="utf-8")

    try:
        write_workbook(workbook_path, rows, summary)
    except Exception as exc:  # noqa: BLE001 - report missing writer deps cleanly.
        print(f"Could not write Excel workbook: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2

    print(f"CSV written to {csv_path}")
    print(f"Summary JSON written to {json_path}")
    print(f"Summary Markdown written to {markdown_path}")
    print(f"Excel workbook written to {workbook_path}")
    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Summarize RAGAS evaluation results.")
    parser.add_argument("--input", required=True, help="Path to ragas_scores.jsonl.")
    parser.add_argument("--output-dir", help="Output directory. Defaults to input parent.")
    parser.add_argument("--workbook", help="Exact output path for evaluation_result.xlsx.")
    return parser.parse_args()


def read_jsonl(path: Path) -> list[dict]:
    rows = []
    with path.open("r", encoding="utf-8") as reader:
        for line in reader:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def flatten_record(record: dict) -> dict:
    dataset = record.get("dataset") or {}
    response = record.get("response") or {}
    contexts = response.get("retrievedContexts") or []
    timings = response.get("timingsMs") or {}
    diagnostics = record.get("diagnostics") or {}
    ragas = record.get("ragas") or {}
    scores = ragas.get("scores") or {}
    errors = ragas.get("errors") or {}

    return {
        "Row ID": record.get("row_id", ""),
        "Language": dataset.get("Language", ""),
        "Question": dataset.get("Question", ""),
        "Generated Answer": response.get("answer", ""),
        "Reference Answer": dataset.get("Answer", ""),
        "Reference Context": dataset.get("Reference Context", ""),
        "Expected File": dataset.get("File Name", ""),
        "Expected Page": dataset.get("Page", ""),
        "Retrieved Files": join_unique(context.get("sourceFile", "") for context in contexts),
        "Retrieved Pages": join_unique(str(context.get("page", "")) for context in contexts),
        "Top Retrieved Context": contexts[0].get("text", "") if contexts else "",
        "Intent": response.get("intent", ""),
        "Detected Language": (response.get("language") or {}).get("language", ""),
        "Retrieval Time Ms": timings.get("retrieval", ""),
        "Generation Time Ms": timings.get("generation", ""),
        "Total Time Ms": timings.get("total", record.get("batch_elapsed_ms", "")),
        "Context Precision": scores.get("context_precision"),
        "Context Recall": scores.get("context_recall"),
        "Context Relevance": scores.get("context_relevance"),
        "Faithfulness": scores.get("faithfulness"),
        "Answer Relevancy": scores.get("answer_relevancy"),
        "Answer Correctness": scores.get("answer_correctness"),
        "File Hit": diagnostics.get("file_hit", False),
        "Page Hit": diagnostics.get("page_hit", False),
        "Expected Rank": diagnostics.get("expected_rank", ""),
        "File Hit@1": diagnostics.get("file_hit_at_1", False),
        "File Hit@3": diagnostics.get("file_hit_at_3", False),
        "File Hit@6": diagnostics.get("file_hit_at_6", False),
        "Page Hit@1": diagnostics.get("page_hit_at_1", False),
        "Page Hit@3": diagnostics.get("page_hit_at_3", False),
        "Page Hit@6": diagnostics.get("page_hit_at_6", False),
        "Paraphrase Number": dataset.get("Paraphrase Number", ""),
        "Original Question": dataset.get("Original Question", ""),
        "Error": record.get("error") or ragas.get("error") or "; ".join(f"{k}: {v}" for k, v in errors.items()),
    }


def join_unique(values: object) -> str:
    seen = []
    for value in values:
        text = str(value).strip()
        if text and text not in seen:
            seen.append(text)
    return " | ".join(seen)


def build_summary(rows: list[dict]) -> dict:
    return {
        "overall": build_group_summary(rows),
        "by_language": summarize_by(rows, "Language"),
        "by_file": summarize_by(rows, "Expected File"),
        "by_original_or_paraphrase": summarize_by(rows, original_or_paraphrase),
        "slowest_20": sorted(rows, key=lambda row: numeric(row.get("Total Time Ms")), reverse=True)[:20],
        "worst_20": sorted(rows, key=average_score)[:20],
    }


def summarize_by(rows: list[dict], key: str | callable) -> list[dict]:
    groups = defaultdict(list)
    for row in rows:
        group_key = key(row) if callable(key) else row.get(key, "")
        groups[group_key or "Unknown"].append(row)

    return [
        {"group": group, **build_group_summary(group_rows)}
        for group, group_rows in sorted(groups.items(), key=lambda item: str(item[0]).lower())
    ]


def original_or_paraphrase(row: dict) -> str:
    return "Paraphrase" if str(row.get("Paraphrase Number", "")).strip() else "Original"


def build_group_summary(rows: list[dict]) -> dict:
    succeeded = [row for row in rows if not row.get("Error")]
    total_times = [numeric(row.get("Total Time Ms")) for row in rows if numeric(row.get("Total Time Ms")) > 0]

    return {
        "total_questions": len(rows),
        "succeeded": len(succeeded),
        "failed": len(rows) - len(succeeded),
        "average_retrieval_time_ms": average_field(rows, "Retrieval Time Ms"),
        "average_generation_time_ms": average_field(rows, "Generation Time Ms"),
        "average_total_time_ms": average_field(rows, "Total Time Ms"),
        "p50_total_time_ms": percentile(total_times, 50),
        "p95_total_time_ms": percentile(total_times, 95),
        "average_context_precision": average_field(rows, "Context Precision"),
        "average_context_recall": average_field(rows, "Context Recall"),
        "average_context_relevance": average_field(rows, "Context Relevance"),
        "average_faithfulness": average_field(rows, "Faithfulness"),
        "average_answer_relevancy": average_field(rows, "Answer Relevancy"),
        "average_answer_correctness": average_field(rows, "Answer Correctness"),
        "file_hit_at_1": ratio(rows, "File Hit@1"),
        "file_hit_at_3": ratio(rows, "File Hit@3"),
        "file_hit_at_6": ratio(rows, "File Hit@6"),
        "page_hit_at_1": ratio(rows, "Page Hit@1"),
        "page_hit_at_3": ratio(rows, "Page Hit@3"),
        "page_hit_at_6": ratio(rows, "Page Hit@6"),
    }


def numeric(value: object) -> float:
    try:
        if value is None or value == "":
            return 0.0
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def average_field(rows: list[dict], field: str) -> float | None:
    values = [numeric(row.get(field)) for row in rows if row.get(field) not in (None, "")]
    return round(statistics.fmean(values), 4) if values else None


def percentile(values: list[float], percentile_value: int) -> float | None:
    if not values:
        return None

    sorted_values = sorted(values)
    index = round((len(sorted_values) - 1) * percentile_value / 100)
    return round(sorted_values[index], 4)


def ratio(rows: list[dict], field: str) -> float | None:
    if not rows:
        return None

    return round(sum(1 for row in rows if bool(row.get(field))) / len(rows), 4)


def average_score(row: dict) -> float:
    values = [
        numeric(row.get(field))
        for field in [
            "Context Precision",
            "Context Recall",
            "Context Relevance",
            "Faithfulness",
            "Answer Relevancy",
            "Answer Correctness",
        ]
        if row.get(field) not in (None, "")
    ]
    return statistics.fmean(values) if values else -1


def write_csv(path: Path, rows: list[dict]) -> None:
    fields = QUESTION_COLUMNS + [
        "File Hit@1",
        "File Hit@3",
        "File Hit@6",
        "Page Hit@1",
        "Page Hit@3",
        "Page Hit@6",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as writer_file:
        writer = csv.DictWriter(writer_file, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_workbook(path: Path, rows: list[dict], summary: dict) -> None:
    from openpyxl import Workbook  # type: ignore[import-not-found]
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    results = workbook.active
    results.title = "Question Results"
    summary_sheet = workbook.create_sheet("Dataset Summary")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    results.append(QUESTION_COLUMNS)
    for cell in results[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row in rows:
        results.append([row.get(column, "") for column in QUESTION_COLUMNS])

    results.freeze_panes = "A2"
    results.auto_filter.ref = results.dimensions
    set_column_widths(results, QUESTION_COLUMNS)

    write_summary_sheet(summary_sheet, summary, header_fill, header_font)
    workbook.save(path)


def set_column_widths(sheet: object, columns: list[str]) -> None:
    width_by_name = {
        "Question": 48,
        "Generated Answer": 56,
        "Reference Answer": 56,
        "Reference Context": 56,
        "Top Retrieved Context": 56,
        "Original Question": 48,
        "Error": 48,
    }
    for index, column in enumerate(columns, start=1):
        width = width_by_name.get(column, min(max(len(column) + 2, 12), 24))
        sheet.column_dimensions[column_letter(index)].width = width

    for row in sheet.iter_rows():
        for cell in row:
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.vertical = "top"
            cell.alignment = alignment


def column_letter(index: int) -> str:
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def write_summary_sheet(sheet: object, summary: dict, header_fill: object, header_font: object) -> None:
    sheet.append(["Metric", "Value"])
    style_header_row(sheet, 1, header_fill, header_font)

    labels = [
        ("Total Questions", "total_questions"),
        ("Succeeded", "succeeded"),
        ("Failed", "failed"),
        ("Average Retrieval Time Ms", "average_retrieval_time_ms"),
        ("Average Generation Time Ms", "average_generation_time_ms"),
        ("Average Total Time Ms", "average_total_time_ms"),
        ("P50 Total Time Ms", "p50_total_time_ms"),
        ("P95 Total Time Ms", "p95_total_time_ms"),
        ("Average Context Precision", "average_context_precision"),
        ("Average Context Recall", "average_context_recall"),
        ("Average Context Relevance", "average_context_relevance"),
        ("Average Faithfulness", "average_faithfulness"),
        ("Average Answer Relevancy", "average_answer_relevancy"),
        ("Average Answer Correctness", "average_answer_correctness"),
        ("File Hit@1", "file_hit_at_1"),
        ("File Hit@3", "file_hit_at_3"),
        ("File Hit@6", "file_hit_at_6"),
        ("Page Hit@1", "page_hit_at_1"),
        ("Page Hit@3", "page_hit_at_3"),
        ("Page Hit@6", "page_hit_at_6"),
    ]

    overall = summary["overall"]
    for label, key in labels:
        sheet.append([label, overall.get(key)])

    current_row = sheet.max_row + 2
    current_row = write_breakdown(sheet, current_row, "Scores by Language", summary["by_language"], header_fill, header_font)
    current_row = write_breakdown(sheet, current_row + 1, "Scores by Source File", summary["by_file"], header_fill, header_font)
    current_row = write_breakdown(
        sheet,
        current_row + 1,
        "Scores by Original/Paraphrase",
        summary["by_original_or_paraphrase"],
        header_fill,
        header_font,
    )
    current_row = write_case_table(sheet, current_row + 1, "Slowest 20 Questions", summary["slowest_20"], header_fill, header_font)
    write_case_table(sheet, current_row + 1, "Worst 20 Questions by Average Score", summary["worst_20"], header_fill, header_font)

    sheet.freeze_panes = "A2"
    for index in range(1, 12):
        sheet.column_dimensions[column_letter(index)].width = 24 if index < 3 else 18
    sheet.column_dimensions["B"].width = 32
    sheet.column_dimensions["C"].width = 48


def style_header_row(sheet: object, row_number: int, header_fill: object, header_font: object) -> None:
    for cell in sheet[row_number]:
        cell.fill = header_fill
        cell.font = header_font
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        alignment.vertical = "top"
        cell.alignment = alignment


def write_breakdown(sheet: object, start_row: int, title: str, rows: list[dict], header_fill: object, header_font: object) -> int:
    sheet.cell(row=start_row, column=1, value=title)
    sheet.cell(row=start_row, column=1).font = FontBold()
    header = [
        "Group",
        "Total",
        "Avg Context Precision",
        "Avg Context Recall",
        "Avg Context Relevance",
        "Avg Faithfulness",
        "Avg Answer Relevancy",
        "Avg Answer Correctness",
        "Avg Total Ms",
        "File Hit@6",
        "Page Hit@6",
    ]
    row_number = start_row + 1
    for col, value in enumerate(header, start=1):
        sheet.cell(row=row_number, column=col, value=value)
    style_header_row(sheet, row_number, header_fill, header_font)

    for item in rows:
        row_number += 1
        values = [
            item.get("group"),
            item.get("total_questions"),
            item.get("average_context_precision"),
            item.get("average_context_recall"),
            item.get("average_context_relevance"),
            item.get("average_faithfulness"),
            item.get("average_answer_relevancy"),
            item.get("average_answer_correctness"),
            item.get("average_total_time_ms"),
            item.get("file_hit_at_6"),
            item.get("page_hit_at_6"),
        ]
        for col, value in enumerate(values, start=1):
            sheet.cell(row=row_number, column=col, value=value)

    return row_number + 1


def write_case_table(sheet: object, start_row: int, title: str, rows: list[dict], header_fill: object, header_font: object) -> int:
    sheet.cell(row=start_row, column=1, value=title)
    sheet.cell(row=start_row, column=1).font = FontBold()
    header = ["Row ID", "Language", "Question", "Total Time Ms", "Average Score", "Expected File", "Expected Page", "File Hit", "Page Hit", "Error"]
    row_number = start_row + 1
    for col, value in enumerate(header, start=1):
        sheet.cell(row=row_number, column=col, value=value)
    style_header_row(sheet, row_number, header_fill, header_font)

    for item in rows:
        row_number += 1
        values = [
            item.get("Row ID"),
            item.get("Language"),
            item.get("Question"),
            item.get("Total Time Ms"),
            round(average_score(item), 4) if average_score(item) >= 0 else None,
            item.get("Expected File"),
            item.get("Expected Page"),
            item.get("File Hit"),
            item.get("Page Hit"),
            item.get("Error"),
        ]
        for col, value in enumerate(values, start=1):
            sheet.cell(row=row_number, column=col, value=value)

    return row_number + 1


def FontBold() -> object:
    from openpyxl.styles import Font  # type: ignore[import-not-found]

    return Font(bold=True)


def build_markdown(summary: dict) -> str:
    overall = summary["overall"]
    lines = [
        "# RAGAS Evaluation Summary",
        "",
        f"- Total questions: {overall['total_questions']}",
        f"- Succeeded: {overall['succeeded']}",
        f"- Failed: {overall['failed']}",
        f"- Average total time ms: {overall['average_total_time_ms']}",
        f"- Average context precision: {overall['average_context_precision']}",
        f"- Average context recall: {overall['average_context_recall']}",
        f"- Average context relevance: {overall['average_context_relevance']}",
        f"- Average faithfulness: {overall['average_faithfulness']}",
        f"- Average answer relevancy: {overall['average_answer_relevancy']}",
        f"- Average answer correctness: {overall['average_answer_correctness']}",
        f"- File hit@6: {overall['file_hit_at_6']}",
        f"- Page hit@6: {overall['page_hit_at_6']}",
        "",
    ]
    return "\n".join(lines)


if __name__ == "__main__":
    raise SystemExit(main())
