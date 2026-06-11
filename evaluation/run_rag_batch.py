#!/usr/bin/env python3
"""Run the policy bot evaluation endpoint for each row in an XLSX dataset."""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
import urllib.error
import urllib.request
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile


DEFAULT_DATASET = Path("evaluation/dataset/test_evaluation_dataset.xlsx")
DEFAULT_CONFIG = Path("evaluation/config.example.json")
REQUIRED_COLUMNS = [
    "Language",
    "Question",
    "Answer",
    "Reference Context",
    "Page",
    "File Name",
    "Original Language",
    "Original Question",
    "Original Answer",
    "Paraphrase Number",
]
STRICT_NS = {"s": "http://purl.oclc.org/ooxml/spreadsheetml/main"}
TRANSITIONAL_NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def main() -> int:
    args = parse_args()
    config = load_json(args.config)
    api_config = config.get("api", {})
    api_base = args.api_base or api_config.get("baseUrl", "http://localhost:5000")
    top_k = args.top_k or int(api_config.get("topK", 6))
    include_prompt = args.include_prompt or bool(api_config.get("includePrompt", False))
    timeout = int(args.timeout or api_config.get("timeoutSeconds", 180))

    dataset_path = Path(args.dataset)
    rows = read_dataset(dataset_path)
    if args.limit is not None:
        rows = rows[: args.limit]

    output_dir = Path(args.output_dir) if args.output_dir else make_output_dir()
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "generated.jsonl"

    completed = read_completed_row_ids(output_path) if args.resume else set()
    endpoint = api_base.rstrip("/") + "/api/evaluation/chat"

    print(f"Dataset: {dataset_path}")
    print(f"Rows: {len(rows)}")
    print(f"Endpoint: {endpoint}")
    print(f"Output: {output_path}")

    with output_path.open("a", encoding="utf-8") as writer:
        for index, row in enumerate(rows, start=1):
            row_id = str(row["Row ID"])
            if row_id in completed:
                continue

            question = str(row.get("Question", "")).strip()
            record = {
                "row_id": row_id,
                "dataset": row,
                "request": {
                    "message": question,
                    "topK": top_k,
                    "includePrompt": include_prompt,
                },
            }

            started = time.perf_counter()
            try:
                response = post_json(endpoint, record["request"], timeout)
                elapsed_ms = int((time.perf_counter() - started) * 1000)
                record["response"] = response
                record["batch_elapsed_ms"] = elapsed_ms
                record["diagnostics"] = build_retrieval_diagnostics(row, response)
                print(f"[{index}/{len(rows)}] row {row_id}: ok")
            except Exception as exc:  # noqa: BLE001 - batch output should capture all failures.
                elapsed_ms = int((time.perf_counter() - started) * 1000)
                record["error"] = f"{type(exc).__name__}: {exc}"
                record["batch_elapsed_ms"] = elapsed_ms
                print(f"[{index}/{len(rows)}] row {row_id}: error: {record['error']}", file=sys.stderr)

            writer.write(json.dumps(record, ensure_ascii=False) + "\n")
            writer.flush()

    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run RAG batch generation for RAGAS evaluation.")
    parser.add_argument("--dataset", default=str(DEFAULT_DATASET), help="Path to XLSX evaluation dataset.")
    parser.add_argument("--config", default=str(DEFAULT_CONFIG), help="Path to evaluation config JSON.")
    parser.add_argument("--api-base", help="API base URL, for example http://localhost:5000.")
    parser.add_argument("--output-dir", help="Output run directory. Defaults to evaluation/outputs/{timestamp}.")
    parser.add_argument("--top-k", type=int, help="Number of retrieved chunks to request.")
    parser.add_argument("--timeout", type=int, help="HTTP timeout in seconds.")
    parser.add_argument("--limit", type=int, help="Only process the first N dataset rows.")
    parser.add_argument("--include-prompt", action="store_true", help="Ask the API to include the generated prompt.")
    parser.add_argument("--no-resume", dest="resume", action="store_false", help="Do not skip existing row IDs.")
    parser.set_defaults(resume=True)
    return parser.parse_args()


def load_json(path: str | Path) -> dict:
    config_path = Path(path)
    if not config_path.exists():
        return {}

    return json.loads(config_path.read_text(encoding="utf-8"))


def make_output_dir() -> Path:
    run_id = datetime.now().strftime("%Y%m%d-%H%M%S")
    return Path("evaluation/outputs") / run_id


def read_dataset(path: Path) -> list[dict]:
    if not path.exists():
        raise FileNotFoundError(path)

    try:
        rows = read_dataset_with_openpyxl(path)
    except Exception:
        rows = read_dataset_from_xlsx_xml(path)

    missing = [column for column in REQUIRED_COLUMNS if column not in rows[0]] if rows else REQUIRED_COLUMNS
    if missing:
        raise ValueError(f"Dataset is missing required columns: {', '.join(missing)}")

    return rows


def read_dataset_with_openpyxl(path: Path) -> list[dict]:
    from openpyxl import load_workbook  # type: ignore[import-not-found]

    workbook = load_workbook(path, read_only=True, data_only=True)
    if not workbook.sheetnames:
        raise ValueError("Workbook has no readable sheets.")

    sheet = workbook[workbook.sheetnames[0]]
    raw_rows = list(sheet.iter_rows(values_only=True))
    return rows_from_matrix(raw_rows)


def read_dataset_from_xlsx_xml(path: Path) -> list[dict]:
    with ZipFile(path) as archive:
        shared_strings = read_shared_strings(archive)
        sheet_names = [name for name in archive.namelist() if name.startswith("xl/worksheets/sheet")]
        if not sheet_names:
            raise ValueError("Workbook has no worksheets.")

        sheet_xml = archive.read(sheet_names[0])
        rows = read_sheet_rows(sheet_xml, shared_strings, STRICT_NS)
        if not rows:
            rows = read_sheet_rows(sheet_xml, shared_strings, TRANSITIONAL_NS)

    return rows_from_matrix(rows)


def read_shared_strings(archive: ZipFile) -> list[str]:
    try:
        root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
    except KeyError:
        return []

    namespace = STRICT_NS if root.tag.startswith("{" + STRICT_NS["s"] + "}") else TRANSITIONAL_NS
    strings = []
    for item in root.findall("s:si", namespace):
        text = "".join(node.text or "" for node in item.findall(".//s:t", namespace))
        strings.append(text)
    return strings


def read_sheet_rows(sheet_xml: bytes, shared_strings: list[str], namespace: dict[str, str]) -> list[list[str]]:
    root = ET.fromstring(sheet_xml)
    rows = []
    for row in root.findall(".//s:sheetData/s:row", namespace):
        values = {}
        for cell in row.findall("s:c", namespace):
            cell_ref = cell.attrib.get("r", "A1")
            index = column_index(cell_ref)
            value_node = cell.find("s:v", namespace)
            value = "" if value_node is None or value_node.text is None else value_node.text
            if cell.attrib.get("t") == "s" and value:
                value = shared_strings[int(value)]
            values[index] = value

        if values:
            max_index = max(values)
            rows.append([values.get(i, "") for i in range(max_index + 1)])
    return rows


def rows_from_matrix(raw_rows: list[tuple | list]) -> list[dict]:
    if not raw_rows:
        return []

    headers = [clean_cell(value) for value in raw_rows[0]]
    rows = []
    for excel_row_number, raw_row in enumerate(raw_rows[1:], start=2):
        if not any(clean_cell(value) for value in raw_row):
            continue

        values = [clean_cell(value) for value in raw_row]
        row = dict(zip(headers, values + [""] * (len(headers) - len(values))))
        row["Row ID"] = str(excel_row_number)
        rows.append(row)
    return rows


def clean_cell(value: object) -> str:
    return "" if value is None else str(value).strip()


def column_index(cell_ref: str) -> int:
    match = re.match(r"([A-Z]+)", cell_ref)
    if not match:
        return 0

    value = 0
    for character in match.group(1):
        value = value * 26 + ord(character) - 64
    return value - 1


def read_completed_row_ids(path: Path) -> set[str]:
    completed = set()
    if not path.exists():
        return completed

    with path.open("r", encoding="utf-8") as reader:
        for line in reader:
            if not line.strip():
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError:
                continue
            completed.add(str(record.get("row_id", "")))
    return completed


def post_json(url: str, payload: dict, timeout: int) -> dict:
    data = json.dumps(payload).encode("utf-8")
    request = urllib.request.Request(
        url,
        data=data,
        method="POST",
        headers={
            "Content-Type": "application/json",
            "Accept": "application/json",
        },
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            body = response.read().decode("utf-8")
            return json.loads(body)
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code}: {detail}") from exc


def build_retrieval_diagnostics(dataset_row: dict, response: dict) -> dict:
    expected_file = str(dataset_row.get("File Name", "")).strip()
    expected_page = parse_int(dataset_row.get("Page"))
    contexts = response.get("retrievedContexts") or []

    expected_rank = None
    file_hit_ranks = []
    page_hit_ranks = []
    for context in contexts:
        rank = parse_int(context.get("rank"))
        source_file = str(context.get("sourceFile", "")).strip()
        page = parse_int(context.get("page"))
        if same_file(source_file, expected_file):
            file_hit_ranks.append(rank)
            if page == expected_page:
                page_hit_ranks.append(rank)
                if expected_rank is None:
                    expected_rank = rank

    return {
        "expected_file": expected_file,
        "expected_page": expected_page,
        "file_hit": bool(file_hit_ranks),
        "page_hit": bool(page_hit_ranks),
        "expected_rank": expected_rank,
        "file_hit_at_1": any(rank <= 1 for rank in file_hit_ranks),
        "file_hit_at_3": any(rank <= 3 for rank in file_hit_ranks),
        "file_hit_at_6": any(rank <= 6 for rank in file_hit_ranks),
        "page_hit_at_1": any(rank <= 1 for rank in page_hit_ranks),
        "page_hit_at_3": any(rank <= 3 for rank in page_hit_ranks),
        "page_hit_at_6": any(rank <= 6 for rank in page_hit_ranks),
    }


def same_file(actual: str, expected: str) -> bool:
    return actual.lower() == expected.lower() or Path(actual).name.lower() == Path(expected).name.lower()


def parse_int(value: object) -> int:
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return 0


if __name__ == "__main__":
    raise SystemExit(main())
